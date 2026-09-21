"""
@Author         : Ailitonia
@Date           : 2026/9/18 14:44
@FileName       : test_002_excel_tools
@Project        : omega-miya
@Description    : Excel 工具单元测试
@GitHub         : https://github.com/Ailitonia
@Software       : PyCharm
"""

import io
import shutil
import warnings
from datetime import UTC, datetime
from typing import TYPE_CHECKING

import pandas as pd
import pytest
from openpyxl import Workbook
from pydantic import BaseModel, ValidationError, create_model

if TYPE_CHECKING:
    from io import BytesIO

    from src.utils.excel_tools import ExcelTools


class SampleModel(BaseModel):
    """全字段必填的混合类型样例模型"""

    name: str
    count: int
    score: float
    enabled: bool
    created_at: datetime


class OptionalModel(BaseModel):
    """含可选字段的样例模型"""

    name: str
    remark: str | None = None
    score: float | None = None


class NullableModel(BaseModel):
    """全可选字段的样例模型"""

    a: int | None = None
    b: str | None = None


MultiHeaderModel = create_model(
    'MultiHeaderModel',
    **{"('a', 'x')": (int | None, None), "('a', 'y')": (int | None, None)},
)
"""多级表头(header=[0, 1])读回列名 str 化后的显式样例模型"""

SAMPLE_MODELS = [
    SampleModel(name='alpha', count=1, score=1.5, enabled=True, created_at=datetime(2024, 1, 1, 12, 30, 45)),
    SampleModel(name='beta', count=-20, score=-0.001, enabled=False, created_at=datetime(2025, 6, 15, 8, 0, 0)),
    SampleModel(
        name='中文测试', count=2147483649, score=1e10, enabled=True, created_at=datetime(2026, 12, 31, 23, 59, 59),
    ),
]


def _make_multi_header_workbook() -> Workbook:
    """构造两个 sheet 的多级表头工作簿, 每个 sheet 前两行为表头"""
    workbook = Workbook()
    sheet1 = workbook.active
    sheet1.title = 'M1'
    sheet1.append(['a', 'a', 'b'])
    sheet1.append(['x', 'y', 'x'])
    sheet1.append([1, 2, 3])
    sheet2 = workbook.create_sheet('M2')
    sheet2.append(['a', 'a'])
    sheet2.append(['x', 'y'])
    sheet2.append([4, 5])
    return workbook


def _read_excel(buf: 'BytesIO', **kwargs) -> 'pd.DataFrame':
    """重置读指针后从内存缓冲区读取 Excel 数据"""
    buf.seek(0)
    return pd.read_excel(buf, **kwargs)


def _excel_file(buf: 'BytesIO') -> 'pd.ExcelFile':
    """重置读指针后从内存缓冲区打开 ExcelFile"""
    buf.seek(0)
    return pd.ExcelFile(buf)


def _cleanup_path_test_folder() -> None:
    """清理路径语义用例产生的临时文件与目录"""
    from src.resource import TemporaryResource

    shutil.rmtree(TemporaryResource('excel_tools_test').path, ignore_errors=True)


# 异常信息片段, 用于 pytest.raises 的 match 校验
_MSG_NOT_FILE = 'is not a file'
_MSG_VALID_STRING = 'valid string'
_MSG_PARSE_FAIL = '解析 Excel 文件失败'
_MSG_WRITE_FAIL = '写入 Excel 文件失败'
_MSG_SHEET_EXISTS = 'already exists'
_MSG_EMPTY_BUFFER = '空缓冲区'
_MSG_BUFFER_UNUSABLE = '缓冲区不可用'


@pytest.fixture
def sample_tools() -> 'ExcelTools[SampleModel]':
    from src.utils.excel_tools import ExcelTools

    return ExcelTools(SampleModel)


@pytest.fixture
def optional_tools() -> 'ExcelTools[OptionalModel]':
    from src.utils.excel_tools import ExcelTools

    return ExcelTools(OptionalModel)


@pytest.fixture
def nullable_tools() -> 'ExcelTools[NullableModel]':
    from src.utils.excel_tools import ExcelTools

    return ExcelTools(NullableModel)


@pytest.fixture
def excel_file() -> 'BytesIO':
    """每个用例独立的内存 Excel 缓冲区, 读写均不落盘"""
    return io.BytesIO()


class TestModuleContract:
    def test_all_exports(self):
        from src.utils import excel_tools

        assert excel_tools.__all__ == ['ExcelFileTarget', 'ExcelTools', 'ExcelToolsException']
        assert excel_tools.ExcelFileTarget is not None
        assert excel_tools.ExcelTools is not None
        assert excel_tools.ExcelToolsException is not None

    def test_init_binds_data_model(self, sample_tools):
        from src.utils.excel_tools import ExcelTools

        assert isinstance(sample_tools, ExcelTools)
        assert sample_tools.data_model is SampleModel

    @pytest.mark.parametrize('invalid_model', [dict, 'not_a_model', 123], ids=['class-dict', 'str', 'int'])
    def test_init_rejects_non_base_model(self, invalid_model):
        """data_model 非 BaseModel 子类时在构造期抛出 TypeError, 不透出裸 AttributeError"""
        from src.utils.excel_tools import ExcelTools

        with pytest.raises(TypeError, match='BaseModel'):
            ExcelTools(invalid_model)


class TestDumpExcel:
    async def test_dump_writes_nonempty_content(self, sample_tools, excel_file):
        await sample_tools.dump_excel(SAMPLE_MODELS, excel_file)

        assert excel_file.getbuffer().nbytes > 0
        assert await sample_tools.load_excel(excel_file) == SAMPLE_MODELS

    async def test_dump_accepts_dicts(self, sample_tools, excel_file):
        data = [m.model_dump() for m in SAMPLE_MODELS]

        await sample_tools.dump_excel(data, excel_file)

        assert await sample_tools.load_excel(excel_file) == SAMPLE_MODELS

    async def test_dump_default_no_index_column(self, sample_tools, excel_file):
        await sample_tools.dump_excel(SAMPLE_MODELS, excel_file)

        columns = _read_excel(excel_file).columns.tolist()
        assert columns == list(SampleModel.model_fields)
        assert 'Unnamed: 0' not in columns

    async def test_dump_with_index_true_keeps_model_loadable(self, sample_tools, excel_file):
        await sample_tools.dump_excel(SAMPLE_MODELS, excel_file, index=True)

        columns = _read_excel(excel_file).columns.tolist()
        assert 'Unnamed: 0' in columns

        assert await sample_tools.load_excel(excel_file) == SAMPLE_MODELS

    async def test_dump_default_sheet_name_is_model_name(self, sample_tools, excel_file):
        await sample_tools.dump_excel(SAMPLE_MODELS, excel_file)

        with _excel_file(excel_file) as excel:
            assert excel.sheet_names == [SampleModel.__name__]

    async def test_dump_custom_sheet_name(self, sample_tools, excel_file):
        await sample_tools.dump_excel(SAMPLE_MODELS, excel_file, sheet_name='Custom')

        with _excel_file(excel_file) as excel:
            assert excel.sheet_names == ['Custom']
        assert await sample_tools.load_excel(excel_file, sheet_name='Custom') == SAMPLE_MODELS

    async def test_dump_creates_missing_parent_dirs(self, sample_tools):
        from src.resource import TemporaryResource

        file = TemporaryResource('excel_tools_test', 'deep', 'nested', 'subdir', 'dump.xlsx')
        try:
            await sample_tools.dump_excel(SAMPLE_MODELS, file)

            assert file.is_file
        finally:
            _cleanup_path_test_folder()

    async def test_dump_empty_data_writes_header_only(self, sample_tools, excel_file):
        await sample_tools.dump_excel([], excel_file)

        assert excel_file.getbuffer().nbytes > 0
        assert _read_excel(excel_file).columns.tolist() == list(SampleModel.model_fields)
        assert await sample_tools.load_excel(excel_file) == []

    async def test_dump_invalid_data_raises_validation_error(self, sample_tools, excel_file):
        with pytest.raises(ValidationError):
            await sample_tools.dump_excel([{'name': 'alpha'}], excel_file)

        assert excel_file.getvalue() == b''

    async def test_dump_overwrites_existing_content(self, sample_tools, excel_file):
        replacement = [SampleModel(
            name='replacement', count=99, score=9.9, enabled=True, created_at=datetime(2024, 2, 2),
        )]

        await sample_tools.dump_excel(SAMPLE_MODELS, excel_file)
        await sample_tools.dump_excel(replacement, excel_file)

        assert await sample_tools.load_excel(excel_file) == replacement

    @pytest.mark.parametrize(
        ('name', 'remark', 'score'),
        [
            ('中文测试', '含"引号"和,逗号', -1.5),
            ('🔐🔑', 'line1\nline2\ttab', 1e10),
            ('a' * 500, None, None),
            ('neg', '-', -0.001),
        ],
        ids=['chinese', 'emoji', 'long-string', 'negative'],
    )
    async def test_dump_special_values_round_trip(self, optional_tools, excel_file, name, remark, score):
        data = [OptionalModel(name=name, remark=remark, score=score)]

        await optional_tools.dump_excel(data, excel_file)

        assert await optional_tools.load_excel(excel_file) == data

    async def test_dump_accepts_generator(self, sample_tools, excel_file):
        """非分块路径同样接受一次性生成器输入(内部全量物化校验)"""
        await sample_tools.dump_excel((m for m in SAMPLE_MODELS), excel_file)

        assert await sample_tools.load_excel(excel_file) == SAMPLE_MODELS

    async def test_dump_subclass_instance_normalized_to_data_model(self, sample_tools, excel_file):
        """data_model 子类实例按 data_model 归一化写出, 子类额外字段不写出"""

        class SubSampleModel(SampleModel):
            extra_field: str = 'extra'

        data = [SubSampleModel(**m.model_dump()) for m in SAMPLE_MODELS]

        await sample_tools.dump_excel(data, excel_file)

        columns = _read_excel(excel_file).columns.tolist()
        assert columns == list(SampleModel.model_fields)
        assert 'extra_field' not in columns
        assert await sample_tools.load_excel(excel_file) == SAMPLE_MODELS


class TestLoadExcel:
    async def test_round_trip(self, sample_tools, excel_file):
        await sample_tools.dump_excel(SAMPLE_MODELS, excel_file)

        assert await sample_tools.load_excel(excel_file) == SAMPLE_MODELS

    async def test_empty_cells_become_none_for_optional_fields(self, optional_tools, excel_file):
        data = [
            OptionalModel(name='a', remark='r', score=1.5),
            OptionalModel(name='b', remark=None, score=None),
        ]

        await optional_tools.dump_excel(data, excel_file)

        assert await optional_tools.load_excel(excel_file) == data

    async def test_empty_cell_on_required_field_raises(self, sample_tools, excel_file):
        pd.DataFrame({
            'name': [None],
            'count': [1],
            'score': [1.5],
            'enabled': [True],
            'created_at': [datetime(2024, 1, 1)],
        }).to_excel(excel_file, index=False)

        with pytest.raises(ValidationError, match=_MSG_VALID_STRING):
            await sample_tools.load_excel(excel_file)

    async def test_load_sheet_by_name_and_index(self, sample_tools, excel_file):
        second = SampleModel(name='second', count=2, score=2.5, enabled=False, created_at=datetime(2024, 2, 2))
        with pd.ExcelWriter(excel_file) as writer:
            pd.DataFrame([SAMPLE_MODELS[0].model_dump()]).to_excel(writer, sheet_name='S1', index=False)
            pd.DataFrame([second.model_dump()]).to_excel(writer, sheet_name='S2', index=False)

        assert await sample_tools.load_excel(excel_file) == [SAMPLE_MODELS[0]]
        assert await sample_tools.load_excel(excel_file, sheet_name='S2') == [second]
        assert await sample_tools.load_excel(excel_file, sheet_name=1) == [second]

    async def test_load_all_sheets_returns_dict(self, sample_tools, excel_file):
        second = SampleModel(name='second', count=2, score=2.5, enabled=False, created_at=datetime(2024, 2, 2))
        with pd.ExcelWriter(excel_file) as writer:
            pd.DataFrame([m.model_dump() for m in SAMPLE_MODELS]).to_excel(writer, sheet_name='S1', index=False)
            pd.DataFrame([second.model_dump()]).to_excel(writer, sheet_name='S2', index=False)

        loaded = await sample_tools.load_excel(excel_file, sheet_name=None)

        assert loaded == {'S1': SAMPLE_MODELS, 'S2': [second]}

    async def test_load_all_sheets_with_multi_header(self, excel_file):
        from src.utils.excel_tools import ExcelTools

        _make_multi_header_workbook().save(excel_file)

        tools = ExcelTools(MultiHeaderModel)
        loaded = await tools.load_excel(excel_file, sheet_name=None, header=[0, 1])

        assert loaded == {
            'M1': [MultiHeaderModel.model_validate({"('a', 'x')": 1, "('a', 'y')": 2})],
            'M2': [MultiHeaderModel.model_validate({"('a', 'x')": 4, "('a', 'y')": 5})],
        }

    async def test_load_with_skiprows(self, sample_tools, excel_file):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['junk title row'])
        sheet.append(list(SampleModel.model_fields))
        for m in SAMPLE_MODELS:
            sheet.append([m.name, m.count, m.score, m.enabled, m.created_at])
        workbook.save(excel_file)

        loaded = await sample_tools.load_excel(excel_file, skiprows=1)

        assert loaded == SAMPLE_MODELS

    async def test_load_with_nrows(self, sample_tools, excel_file):
        await sample_tools.dump_excel(SAMPLE_MODELS, excel_file)

        assert await sample_tools.load_excel(excel_file, nrows=2) == SAMPLE_MODELS[:2]

    async def test_load_nonexistent_file_raises(self, sample_tools):
        from src.resource import ResourceNotFileError, TemporaryResource

        with pytest.raises(ResourceNotFileError, match=_MSG_NOT_FILE):
            await sample_tools.load_excel(TemporaryResource('excel_tools_test', 'not_exists.xlsx'))

    async def test_load_directory_raises(self, sample_tools):
        from src.resource import ResourceNotFileError, TemporaryResource

        folder = TemporaryResource('excel_tools_test', 'as_load_dir')
        folder.path.mkdir(parents=True, exist_ok=True)
        try:
            with pytest.raises(ResourceNotFileError, match=_MSG_NOT_FILE):
                await sample_tools.load_excel(folder)
        finally:
            _cleanup_path_test_folder()

    async def test_load_header_only_returns_empty(self, sample_tools, excel_file):
        await sample_tools.dump_excel([], excel_file)

        assert await sample_tools.load_excel(excel_file) == []

    async def test_fully_empty_rows_become_none_models(self, nullable_tools, excel_file):
        pd.DataFrame({'a': [1, None, 3], 'b': ['x', None, 'z']}).to_excel(excel_file, index=False)

        loaded = await nullable_tools.load_excel(excel_file)

        assert loaded == [
            NullableModel(a=1, b='x'),
            NullableModel(a=None, b=None),
            NullableModel(a=3, b='z'),
        ]

    async def test_load_with_nrows_zero(self, sample_tools, excel_file):
        await sample_tools.dump_excel(SAMPLE_MODELS, excel_file)

        assert await sample_tools.load_excel(excel_file, nrows=0) == []

    async def test_load_truly_empty_sheet_returns_empty(self, sample_tools, excel_file):
        """仅含从未写入的空 sheet 的工作簿读取返回空列表(0 列边界)"""
        Workbook().save(excel_file)

        assert await sample_tools.load_excel(excel_file) == []


class TestAutoLoadExcel:
    async def test_generates_fields_from_header(self, excel_file):
        from src.utils.excel_tools import ExcelTools

        pd.DataFrame({'a': [1, 2], 'b': [1.5, 2.5], 'c': ['x', 'y']}).to_excel(excel_file, index=False)

        loaded = await ExcelTools.auto_load_excel(excel_file)

        assert list(type(loaded[0]).model_fields) == ['a', 'b', 'c']
        assert [m.model_dump() for m in loaded] == [
            {'a': '1', 'b': '1.5', 'c': 'x'},
            {'a': '2', 'b': '2.5', 'c': 'y'},
        ]

    async def test_empty_cells_become_none(self, excel_file):
        from src.utils.excel_tools import ExcelTools

        pd.DataFrame({'a': [1, None], 'b': ['x', 'y']}).to_excel(excel_file, index=False)

        loaded = await ExcelTools.auto_load_excel(excel_file)

        assert [m.model_dump() for m in loaded] == [
            {'a': '1.0', 'b': 'x'},
            {'a': None, 'b': 'y'},
        ]

    async def test_empty_cells_never_become_nan_string(self, excel_file):
        from src.utils.excel_tools import ExcelTools

        pd.DataFrame({'a': [None, None], 'b': [None, 2]}).to_excel(excel_file, index=False)

        loaded = await ExcelTools.auto_load_excel(excel_file)

        for m in loaded:
            assert 'nan' not in m.model_dump().values()
        assert loaded[1].model_dump() == {'a': None, 'b': '2.0'}

    async def test_date_column_stringified(self, excel_file):
        from src.utils.excel_tools import ExcelTools

        pd.DataFrame({'d': [datetime(2024, 1, 1, 12, 30)]}).to_excel(excel_file, index=False)

        loaded = await ExcelTools.auto_load_excel(excel_file)

        assert [m.model_dump() for m in loaded] == [{'d': '2024-01-01 12:30:00'}]

    async def test_bool_column_stringified(self, excel_file):
        from src.utils.excel_tools import ExcelTools

        pd.DataFrame({'flag': [True, False]}).to_excel(excel_file, index=False)

        loaded = await ExcelTools.auto_load_excel(excel_file)

        assert [m.model_dump() for m in loaded] == [{'flag': 'True'}, {'flag': 'False'}]

    async def test_duplicate_and_blank_headers(self, excel_file):
        from src.utils.excel_tools import ExcelTools

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['a', 'a', ''])
        sheet.append([1, 2, 3])
        workbook.save(excel_file)

        loaded = await ExcelTools.auto_load_excel(excel_file)

        assert list(type(loaded[0]).model_fields) == ['a', 'a.1', 'Unnamed: 2']
        assert loaded[0].model_dump() == {'a': '1', 'a.1': '2', 'Unnamed: 2': '3'}

    async def test_round_trip_from_dump(self, sample_tools, excel_file):
        from src.utils.excel_tools import ExcelTools

        await sample_tools.dump_excel(SAMPLE_MODELS, excel_file)

        loaded = await ExcelTools.auto_load_excel(excel_file)

        assert [m.model_dump() for m in loaded] == [
            {
                'name': 'alpha',
                'count': '1',
                'score': '1.5',
                'enabled': 'True',
                'created_at': '2024-01-01 12:30:45',
            },
            {
                'name': 'beta',
                'count': '-20',
                'score': '-0.001',
                'enabled': 'False',
                'created_at': '2025-06-15 08:00:00',
            },
            {
                'name': '中文测试',
                'count': '2147483649',
                'score': '10000000000.0',
                'enabled': 'True',
                'created_at': '2026-12-31 23:59:59',
            },
        ]

    async def test_nonexistent_file_raises(self):
        from src.resource import ResourceNotFileError, TemporaryResource
        from src.utils.excel_tools import ExcelTools

        with pytest.raises(ResourceNotFileError, match=_MSG_NOT_FILE):
            await ExcelTools.auto_load_excel(TemporaryResource('excel_tools_test', 'auto_not_exists.xlsx'))

    async def test_header_none_generates_integer_fields(self, excel_file):
        from src.utils.excel_tools import ExcelTools

        pd.DataFrame({'x': [1, 2], 'y': [3, 4]}).to_excel(excel_file, index=False, header=False)

        loaded = await ExcelTools.auto_load_excel(excel_file, header=None)

        assert list(type(loaded[0]).model_fields) == ['0', '1']
        assert [m.model_dump() for m in loaded] == [{'0': '1', '1': '3'}, {'0': '2', '1': '4'}]

    async def test_generated_model_fields_are_optional(self, excel_file):
        from src.utils.excel_tools import ExcelTools

        pd.DataFrame({'a': [1], 'b': ['x']}).to_excel(excel_file, index=False)

        loaded = await ExcelTools.auto_load_excel(excel_file)

        dynamic_model = type(loaded[0])
        assert all(not field.is_required() for field in dynamic_model.model_fields.values())
        assert dynamic_model(a=None, b=None).model_dump() == {'a': None, 'b': None}

    async def test_auto_load_all_sheets_returns_dict_with_independent_models(self, excel_file):
        from src.utils.excel_tools import ExcelTools

        with pd.ExcelWriter(excel_file) as writer:
            pd.DataFrame({'x': [1, 2]}).to_excel(writer, sheet_name='A', index=False)
            pd.DataFrame({'y': [3]}).to_excel(writer, sheet_name='B', index=False)

        loaded = await ExcelTools.auto_load_excel(excel_file, sheet_name=None)

        assert set(loaded) == {'A', 'B'}
        assert list(type(loaded['A'][0]).model_fields) == ['x']
        assert list(type(loaded['B'][0]).model_fields) == ['y']
        assert [m.model_dump() for m in loaded['A']] == [{'x': '1'}, {'x': '2'}]
        assert [m.model_dump() for m in loaded['B']] == [{'y': '3'}]

    async def test_auto_load_multi_header_stringified_tuple_fields(self, excel_file):
        from src.utils.excel_tools import ExcelTools

        _make_multi_header_workbook().save(excel_file)

        loaded = await ExcelTools.auto_load_excel(excel_file, header=[0, 1])

        assert list(type(loaded[0]).model_fields) == ["('a', 'x')", "('a', 'y')", "('b', 'x')"]
        assert [m.model_dump() for m in loaded] == [
            {"('a', 'x')": '1', "('a', 'y')": '2', "('b', 'x')": '3'},
        ]

    async def test_fully_empty_rows_become_all_none_models(self, excel_file):
        from src.utils.excel_tools import ExcelTools

        pd.DataFrame({'a': [1, None, 3], 'b': ['x', None, 'z']}).to_excel(excel_file, index=False)

        loaded = await ExcelTools.auto_load_excel(excel_file)

        assert [m.model_dump() for m in loaded] == [
            {'a': '1.0', 'b': 'x'},
            {'a': None, 'b': None},
            {'a': '3.0', 'b': 'z'},
        ]

    async def test_truly_empty_sheet_returns_empty(self, excel_file):
        """空 sheet(0 列)自动建模返回空列表"""
        from src.utils.excel_tools import ExcelTools

        Workbook().save(excel_file)

        assert await ExcelTools.auto_load_excel(excel_file) == []

    async def test_model_config_column_renamed_and_preserved(self, excel_file):
        """model_config 列被重命名建模, 原始列数据不丢失"""
        from src.utils.excel_tools import ExcelTools

        pd.DataFrame({'model_config': [1], 'a': [2]}).to_excel(excel_file, index=False)

        loaded = await ExcelTools.auto_load_excel(excel_file)

        assert list(type(loaded[0]).model_fields) == ['col_model_config', 'a']
        assert loaded[0].model_dump() == {'col_model_config': '1', 'a': '2'}

    async def test_base_model_member_columns_renamed_and_preserved(self, excel_file):
        """与 BaseModel 成员冲突及 model_/下划线开头的列名重命名后数据保留, 且无 pydantic 告警"""
        from src.utils.excel_tools import ExcelTools

        pd.DataFrame({
            'model_dump': [1],
            'model_validate': [2],
            'json': [3],
            'schema': [4],
            '_secret': [5],
            'x': [6],
        }).to_excel(excel_file, index=False)

        with warnings.catch_warnings(record=True) as warning_record:
            warnings.simplefilter('always')
            loaded = await ExcelTools.auto_load_excel(excel_file)

        assert not [w for w in warning_record if issubclass(w.category, UserWarning)]
        assert list(type(loaded[0]).model_fields) == [
            'col_model_dump', 'col_model_validate', 'col_json', 'col_schema', 'col__secret', 'x',
        ]
        assert loaded[0].model_dump() == {
            'col_model_dump': '1',
            'col_model_validate': '2',
            'col_json': '3',
            'col_schema': '4',
            'col__secret': '5',
            'x': '6',
        }

    async def test_rename_candidate_collision_deduplicated(self, excel_file):
        """重命名候选与真实列名碰撞时追加序号去重, 两列数据均保留"""
        from src.utils.excel_tools import ExcelTools

        pd.DataFrame({'model_config': [1], 'col_model_config': [2]}).to_excel(excel_file, index=False)

        loaded = await ExcelTools.auto_load_excel(excel_file)

        assert list(type(loaded[0]).model_fields) == ['col_model_config_2', 'col_model_config']
        assert loaded[0].model_dump() == {'col_model_config_2': '1', 'col_model_config': '2'}


class TestRoundTrip:
    async def test_multiple_round_trips_preserve_data(self, sample_tools, excel_file):
        await sample_tools.dump_excel(SAMPLE_MODELS, excel_file)
        first = await sample_tools.load_excel(excel_file)
        await sample_tools.dump_excel(first, excel_file)
        second = await sample_tools.load_excel(excel_file)

        assert first == SAMPLE_MODELS
        assert second == SAMPLE_MODELS

    async def test_round_trip_with_none_optionals(self, optional_tools, excel_file):
        data = [OptionalModel(name='a', remark='r', score=None), OptionalModel(name='b', remark=None, score=2.5)]

        await optional_tools.dump_excel(data, excel_file)

        assert await optional_tools.load_excel(excel_file) == data

    async def test_datetime_microseconds_truncated_to_milliseconds(self, sample_tools, excel_file):
        """微秒级 datetime 往返被截断到毫秒(Excel 序列日期固有限制, 已知行为快照)"""
        data = [SampleModel(
            name='ms', count=1, score=1.0, enabled=True, created_at=datetime(2024, 1, 1, 12, 30, 45, 123456),
        )]

        await sample_tools.dump_excel(data, excel_file)
        loaded = await sample_tools.load_excel(excel_file)

        assert loaded[0].created_at == datetime(2024, 1, 1, 12, 30, 45, 123000)
        assert loaded != data

    async def test_dump_failure_leaves_buffer_untouched(self, sample_tools, excel_file):
        """非分块写出中途失败时缓冲区目标保持写入前内容(临时目标原子写入)"""
        from src.utils.excel_tools import ExcelToolsException

        await sample_tools.dump_excel(SAMPLE_MODELS, excel_file)
        original_content = excel_file.getvalue()

        bad_data = [SampleModel(
            name='tz', count=1, score=1.0, enabled=True, created_at=datetime(2024, 1, 1, tzinfo=UTC),
        )]
        with pytest.raises(ExcelToolsException, match=_MSG_WRITE_FAIL):
            await sample_tools.dump_excel(bad_data, excel_file)

        assert excel_file.getvalue() == original_content
        assert await sample_tools.load_excel(excel_file) == SAMPLE_MODELS


class TestAppendExcel:
    async def test_append_new_sheet_preserves_existing(self, sample_tools, excel_file):
        extra = [SampleModel(name='extra', count=9, score=9.5, enabled=True, created_at=datetime(2024, 9, 9))]

        await sample_tools.dump_excel(SAMPLE_MODELS, excel_file)
        await sample_tools.append_excel(extra, excel_file, sheet_name='Extra')

        with _excel_file(excel_file) as excel:
            assert excel.sheet_names == [SampleModel.__name__, 'Extra']
        assert await sample_tools.load_excel(excel_file) == SAMPLE_MODELS
        assert await sample_tools.load_excel(excel_file, sheet_name='Extra') == extra

    async def test_append_default_sheet_name_is_model_name(self, sample_tools, excel_file):
        pd.DataFrame({'old': [1]}).to_excel(excel_file, sheet_name='Old', index=False)

        await sample_tools.append_excel(SAMPLE_MODELS, excel_file)

        with _excel_file(excel_file) as excel:
            assert excel.sheet_names == ['Old', SampleModel.__name__]

    async def test_append_nonexistent_file_raises(self, sample_tools):
        from src.resource import ResourceNotFileError, TemporaryResource

        with pytest.raises(ResourceNotFileError, match=_MSG_NOT_FILE):
            await sample_tools.append_excel(
                SAMPLE_MODELS, TemporaryResource('excel_tools_test', 'append_not_exists.xlsx')
            )

    async def test_append_duplicate_sheet_raises_by_default(self, sample_tools, excel_file):
        from src.utils.excel_tools import ExcelToolsException

        await sample_tools.dump_excel(SAMPLE_MODELS, excel_file)

        with pytest.raises(ExcelToolsException, match=_MSG_SHEET_EXISTS):
            await sample_tools.append_excel(SAMPLE_MODELS, excel_file)

    async def test_append_replace_existing_sheet(self, sample_tools, excel_file):
        replacement = [SampleModel(
            name='replacement', count=99, score=9.9, enabled=True, created_at=datetime(2024, 2, 2),
        )]

        await sample_tools.dump_excel(SAMPLE_MODELS, excel_file)
        await sample_tools.append_excel(replacement, excel_file, if_sheet_exists='replace')

        with _excel_file(excel_file) as excel:
            assert excel.sheet_names == [SampleModel.__name__]
        assert await sample_tools.load_excel(excel_file) == replacement

    async def test_append_replace_creates_sheet_when_missing(self, sample_tools, excel_file):
        pd.DataFrame({'old': [1]}).to_excel(excel_file, sheet_name='Old', index=False)

        await sample_tools.append_excel(SAMPLE_MODELS, excel_file, sheet_name='New', if_sheet_exists='replace')

        with _excel_file(excel_file) as excel:
            assert excel.sheet_names == ['Old', 'New']
        assert await sample_tools.load_excel(excel_file, sheet_name='New') == SAMPLE_MODELS

    async def test_append_empty_data_writes_header_only_sheet(self, sample_tools, excel_file):
        pd.DataFrame({'old': [1]}).to_excel(excel_file, sheet_name='Old', index=False)

        await sample_tools.append_excel([], excel_file, sheet_name='Extra')

        assert await sample_tools.load_excel(excel_file, sheet_name='Extra') == []
        assert _read_excel(excel_file, sheet_name='Extra').columns.tolist() == list(SampleModel.model_fields)

    async def test_append_to_empty_buffer_raises(self, sample_tools, excel_file):
        from src.utils.excel_tools import ExcelToolsException

        with pytest.raises(ExcelToolsException, match=_MSG_EMPTY_BUFFER):
            await sample_tools.append_excel(SAMPLE_MODELS, excel_file)

    async def test_append_to_garbage_buffer_wrapped(self, sample_tools, excel_file):
        from src.utils.excel_tools import ExcelToolsException

        excel_file.write(b'this is not an excel file')

        with pytest.raises(ExcelToolsException, match=_MSG_WRITE_FAIL):
            await sample_tools.append_excel(SAMPLE_MODELS, excel_file)

    async def test_append_with_index_writes_unnamed_column(self, sample_tools, excel_file):
        """append 配合 index=True 时追加 sheet 含索引列, 不影响默认 load(extra 列被忽略)"""
        pd.DataFrame({'old': [1]}).to_excel(excel_file, sheet_name='Old', index=False)

        await sample_tools.append_excel(SAMPLE_MODELS, excel_file, sheet_name='New', index=True)

        columns = _read_excel(excel_file, sheet_name='New').columns.tolist()
        assert columns == ['Unnamed: 0', *SampleModel.model_fields]
        assert await sample_tools.load_excel(excel_file, sheet_name='New') == SAMPLE_MODELS


class TestSheetNameValidation:
    @pytest.mark.parametrize(
        'sheet_name',
        ['', 'x' * 32, 'a/b', 'a*b'],
        ids=['empty', 'too-long', 'invalid-slash', 'invalid-asterisk'],
    )
    async def test_dump_invalid_sheet_name_raises(self, sample_tools, excel_file, sheet_name):
        """非法 sheet 名在写入前校验失败, 目标不被触碰"""
        from src.utils.excel_tools import ExcelToolsException

        with pytest.raises(ExcelToolsException, match='sheet 名'):
            await sample_tools.dump_excel(SAMPLE_MODELS, excel_file, sheet_name=sheet_name)

        assert excel_file.getvalue() == b''

    @pytest.mark.parametrize(
        'sheet_name',
        ['', 'x' * 32, 'a/b', 'a[b'],
        ids=['empty', 'too-long', 'invalid-slash', 'invalid-bracket'],
    )
    async def test_append_invalid_sheet_name_raises(self, sample_tools, excel_file, sheet_name):
        from src.utils.excel_tools import ExcelToolsException

        pd.DataFrame({'old': [1]}).to_excel(excel_file, sheet_name='Old', index=False)

        with pytest.raises(ExcelToolsException, match='sheet 名'):
            await sample_tools.append_excel(SAMPLE_MODELS, excel_file, sheet_name=sheet_name)

        with _excel_file(excel_file) as excel:
            assert excel.sheet_names == ['Old']

    async def test_max_length_sheet_name_allowed(self, sample_tools, excel_file):
        """31 字符上限内的 sheet 名合法"""
        await sample_tools.dump_excel(SAMPLE_MODELS, excel_file, sheet_name='x' * 31)

        with _excel_file(excel_file) as excel:
            assert excel.sheet_names == ['x' * 31]


class TestDumpExcelInChunks:
    async def test_in_chunks_equals_full_dump(self, sample_tools):
        chunked = io.BytesIO()
        full = io.BytesIO()

        await sample_tools.dump_excel((m for m in SAMPLE_MODELS), chunked, in_chunks=True)
        await sample_tools.dump_excel(SAMPLE_MODELS, full)

        assert await sample_tools.load_excel(chunked) == SAMPLE_MODELS
        assert await sample_tools.load_excel(chunked) == await sample_tools.load_excel(full)

    async def test_in_chunks_default_off(self, sample_tools, excel_file):
        await sample_tools.dump_excel(SAMPLE_MODELS, excel_file)

        assert await sample_tools.load_excel(excel_file) == SAMPLE_MODELS

    async def test_in_chunks_empty_data_writes_header_only(self, sample_tools, excel_file):
        await sample_tools.dump_excel([], excel_file, in_chunks=True)

        assert excel_file.getbuffer().nbytes > 0
        assert _read_excel(excel_file).columns.tolist() == list(SampleModel.model_fields)
        assert await sample_tools.load_excel(excel_file) == []

    async def test_in_chunks_custom_sheet_name_single_sheet(self, sample_tools, excel_file):
        await sample_tools.dump_excel(SAMPLE_MODELS, excel_file, sheet_name='Chunked', in_chunks=True)

        with _excel_file(excel_file) as excel:
            assert excel.sheet_names == ['Chunked']
        assert await sample_tools.load_excel(excel_file, sheet_name='Chunked') == SAMPLE_MODELS

    async def test_in_chunks_accepts_generator_input(self, sample_tools, excel_file):
        def _gen():
            yield from SAMPLE_MODELS

        await sample_tools.dump_excel(_gen(), excel_file, in_chunks=True)

        assert await sample_tools.load_excel(excel_file) == SAMPLE_MODELS

    async def test_in_chunks_with_index_writes_zero_index_column(self, sample_tools, excel_file):
        """index=True 时逐行写出的索引列值均为 0, 数据列不受影响(固有行为快照)"""
        await sample_tools.dump_excel(SAMPLE_MODELS, excel_file, index=True, in_chunks=True)

        df = _read_excel(excel_file)
        assert 'Unnamed: 0' in df.columns
        assert df['Unnamed: 0'].tolist() == [0] * len(SAMPLE_MODELS)
        loaded = await sample_tools.load_excel(excel_file)
        assert loaded == SAMPLE_MODELS

    async def test_in_chunks_first_item_invalid_raises_validation_error(self, sample_tools, excel_file):
        """首条数据校验失败时裸抛 ValidationError, 不被包装成 ExcelToolsException"""
        from src.utils.excel_tools import ExcelToolsException

        with pytest.raises(ValidationError) as exc_info:
            await sample_tools.dump_excel([{'name': 'alpha'}], excel_file, in_chunks=True)

        assert not isinstance(exc_info.value, ExcelToolsException)
        assert excel_file.getvalue() == b''

    async def test_in_chunks_later_item_invalid_raises_validation_error(self, sample_tools, excel_file):
        """后续条目校验失败时同样裸抛 ValidationError, 与非分块路径语义一致"""
        from src.utils.excel_tools import ExcelToolsException

        with pytest.raises(ValidationError) as exc_info:
            await sample_tools.dump_excel(
                [SAMPLE_MODELS[0], {'name': 'bad'}], excel_file, in_chunks=True
            )

        assert not isinstance(exc_info.value, ExcelToolsException)

    async def test_in_chunks_later_item_invalid_leaves_buffer_untouched(self, sample_tools, excel_file):
        """后续条目校验失败时缓冲区目标保持写入前内容(临时目标原子写入)"""
        await sample_tools.dump_excel(SAMPLE_MODELS, excel_file)
        original_content = excel_file.getvalue()

        with pytest.raises(ValidationError):
            await sample_tools.dump_excel(
                [SAMPLE_MODELS[0], {'name': 'bad'}], excel_file, in_chunks=True
            )

        assert excel_file.getvalue() == original_content
        assert await sample_tools.load_excel(excel_file) == SAMPLE_MODELS

    async def test_in_chunks_later_item_invalid_leaves_file_untouched(self, sample_tools):
        """后续条目校验失败时文件目标保持写入前内容, 且不残留临时文件"""
        from src.resource import TemporaryResource

        file = TemporaryResource('excel_tools_test', 'atomic_chunks.xlsx')
        try:
            await sample_tools.dump_excel(SAMPLE_MODELS, file)

            with pytest.raises(ValidationError):
                await sample_tools.dump_excel(
                    [SAMPLE_MODELS[0], {'name': 'bad'}], file, in_chunks=True
                )

            assert await sample_tools.load_excel(file) == SAMPLE_MODELS
            assert [p for p in file.path.parent.iterdir() if p.name != file.path.name] == []
        finally:
            _cleanup_path_test_folder()

    async def test_write_excel_in_chunks_empty_iterator(self, excel_file):
        """空迭代器直接调用分块写入, 产出仅含空 sheet 的合法文件(私有方法健壮性)"""
        from src.utils.excel_tools import ExcelTools

        ExcelTools._write_excel_in_chunks(iter([]), excel_file, index=False, sheet_name='Empty')

        with _excel_file(excel_file) as excel:
            assert excel.sheet_names == ['Empty']
        assert _read_excel(excel_file).empty

    async def test_in_chunks_creates_missing_parent_dirs(self, sample_tools):
        from src.resource import TemporaryResource

        file = TemporaryResource('excel_tools_test', 'deep', 'nested', 'dump_in_chunks.xlsx')
        try:
            await sample_tools.dump_excel(SAMPLE_MODELS, file, in_chunks=True)

            assert file.is_file
            assert await sample_tools.load_excel(file) == SAMPLE_MODELS
        finally:
            _cleanup_path_test_folder()


class TestExceptionWrapping:
    async def test_read_corrupt_buffer_wrapped(self, sample_tools, excel_file):
        from src.utils.excel_tools import ExcelToolsException

        excel_file.write(b'this is not an excel file')

        with pytest.raises(ExcelToolsException, match=_MSG_PARSE_FAIL) as exc_info:
            await sample_tools.load_excel(excel_file)

        assert exc_info.value.__cause__ is not None
        assert not isinstance(exc_info.value.__cause__, ExcelToolsException)

    async def test_read_missing_sheet_wrapped_with_sheet_name(self, sample_tools, excel_file):
        from src.utils.excel_tools import ExcelToolsException

        await sample_tools.dump_excel(SAMPLE_MODELS, excel_file)

        with pytest.raises(ExcelToolsException, match='NOPE'):
            await sample_tools.load_excel(excel_file, sheet_name='NOPE')

    async def test_write_to_directory_wrapped(self, sample_tools):
        from src.resource import TemporaryResource
        from src.utils.excel_tools import ExcelToolsException

        target = TemporaryResource('excel_tools_test', 'as_dir.xlsx')
        target.path.mkdir(parents=True, exist_ok=True)
        try:
            with pytest.raises(ExcelToolsException, match=_MSG_WRITE_FAIL):
                await sample_tools.dump_excel(SAMPLE_MODELS, target)
        finally:
            _cleanup_path_test_folder()

    async def test_exception_contract(self):
        from src.exception import OmegaException
        from src.utils.excel_tools import ExcelToolsException

        assert issubclass(ExcelToolsException, OmegaException)

        exc = ExcelToolsException('some message')
        assert exc.message == 'some message'
        assert 'some message' in str(exc)
        assert repr(exc) == "ExcelToolsException(message='some message')"

    async def test_validation_error_not_wrapped(self, sample_tools, excel_file):
        from src.utils.excel_tools import ExcelToolsException

        with pytest.raises(ValidationError) as exc_info:
            await sample_tools.dump_excel([{'name': 'alpha'}], excel_file)

        assert not isinstance(exc_info.value, ExcelToolsException)
        assert excel_file.getvalue() == b''

    async def test_read_closed_buffer_wrapped(self, sample_tools):
        """已关闭缓冲区在目标检查阶段抛出 ExcelToolsException, 不透出裸 ValueError"""
        from src.utils.excel_tools import ExcelToolsException

        buffer = io.BytesIO()
        buffer.close()

        with pytest.raises(ExcelToolsException, match=_MSG_BUFFER_UNUSABLE):
            await sample_tools.load_excel(buffer)

    async def test_dump_closed_buffer_wrapped(self, sample_tools):
        from src.utils.excel_tools import ExcelToolsException

        buffer = io.BytesIO()
        buffer.close()

        with pytest.raises(ExcelToolsException, match=_MSG_BUFFER_UNUSABLE):
            await sample_tools.dump_excel(SAMPLE_MODELS, buffer)

    async def test_append_closed_buffer_wrapped(self, sample_tools):
        from src.utils.excel_tools import ExcelToolsException

        buffer = io.BytesIO()
        buffer.close()

        with pytest.raises(ExcelToolsException, match=_MSG_BUFFER_UNUSABLE):
            await sample_tools.append_excel(SAMPLE_MODELS, buffer)

    async def test_read_non_seekable_stream_wrapped(self, sample_tools):
        """不可 seek 的流在目标检查阶段抛出 ExcelToolsException, 不透出裸 UnsupportedOperation"""
        from src.utils.excel_tools import ExcelToolsException

        class NonSeekableStream:
            def read(self, *args):
                return b''

            def seek(self, *args):
                raise io.UnsupportedOperation('not seekable')

        with pytest.raises(ExcelToolsException, match=_MSG_BUFFER_UNUSABLE):
            await sample_tools.load_excel(NonSeekableStream())

    async def test_dump_timezone_aware_datetime_wrapped(self, sample_tools, excel_file):
        """时区感知 datetime 无法写出(openpyxl 限制), 包装为 ExcelToolsException"""
        from src.utils.excel_tools import ExcelToolsException

        data = [SampleModel(
            name='tz', count=1, score=1.0, enabled=True, created_at=datetime(2024, 1, 1, tzinfo=UTC),
        )]

        with pytest.raises(ExcelToolsException, match=_MSG_WRITE_FAIL):
            await sample_tools.dump_excel(data, excel_file)
